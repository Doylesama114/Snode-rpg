
# -*- coding: utf-8 -*-
import sys, json, re, os, glob, collections
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

ALIAS = {'限制':'施展限制','费用':'标识','疲劳值消耗':'疲劳消耗','疲劳值耗':'疲劳消耗','展条件':'施展条件'}
BASE = ['前置条件','额外条件','施展时间','施展距离','持续时间','疲劳消耗','关键词','施展条件','施展限制','标识','描述']
RECIPE = ['研发材料','研发时间','参考价格','负重','配方材料','使用条件','使用限制','类别','制作时间','配方','效果','注释']
ALL = BASE + RECIPE
KEYS = sorted(set(ALL) | set(ALIAS), key=len, reverse=True)
LABEL_RE = re.compile(r'^\s*(' + '|'.join(KEYS) + r')\s*[：:]\s*')

def iter_body(parent):
    for child in parent.element.body.iterchildren():
        if child.tag == qn('w:p'): yield ('p', Paragraph(child, parent))
        elif child.tag == qn('w:tbl'): yield ('t', Table(child, parent))
def rows_of(t):
    out=[]
    for row in t.rows:
        seen=set(); cells=[]
        for c in row.cells:
            if id(c._tc) in seen: continue
            seen.add(id(c._tc)); cells.append(c)
        out.append(cells)
    return out
def ct(c): return ''.join(p.text for p in c.paragraphs).strip()
def marks_of(c):
    raw=''.join(p.text for p in c.paragraphs); dots=raw.count('●'); colors=[]
    for p in c.paragraphs:
        for r in p.runs:
            if '●' in r.text:
                try: col=r.font.color.rgb
                except Exception: col=None
                colors.extend([str(col) if col else None]*r.text.count('●'))
    return dots, colors
def is_name_list(rows):
    for cells in rows:
        ne=[c for c in cells if ct(c)]
        if len(ne)!=1: return False
        t=ct(ne[0])
        if not t or len(t)>24 or '：' in t or ':' in t or '。' in t: return False
    return True

def survey(path):
    doc=Document(path); style=None; tier=None
    st=dict(doc=os.path.basename(path), skills=[], unknown=collections.Counter())
    for kind,b in iter_body(doc):
        if kind=='p':
            t=b.text.strip()
            if not t: continue
            if re.match(r'^[一二三四五六七八九十]+阶天赋树$',t): tier=t
            elif t.endswith('风格') and len(t)<=12 and '天赋树' not in t: style=t; tier=None
            elif t in ('起始特性','初始专长'): style='起始'; tier=None
            continue
        rows=rows_of(b)
        if not rows or is_name_list(rows): continue
        if not any(LABEL_RE.match(ct(c)) for cells in rows for c in cells): continue
        name=ct(rows[0][0]).strip()
        if not name or len(name)>30: continue
        sk=dict(name=name, style=style, tier=tier, fields={}, marks=[], dots=0, colors=[])
        cur=None
        for cells in rows[1:]:
            for c in cells:
                txt=ct(c)
                if not txt: continue
                m=LABEL_RE.match(txt)
                if m:
                    lab=ALIAS.get(m.group(1), m.group(1)); rest=txt[m.end():].strip(); cur=lab
                    if lab=='标识':
                        dots,colors=marks_of(c)
                        sk['marks'].append(dict(dots=dots,colors=colors))
                        sk['dots']+=dots; sk['colors']+= [x for x in colors if x]
                        sk['fields'].setdefault('标识',[]).append('●'*dots)
                    else:
                        sk['fields'].setdefault(lab,[]).append(rest)
                    continue
                if cur=='描述': sk['fields'].setdefault('描述',[]).append(txt)
                elif len(txt)>12: cur='描述'; sk['fields'].setdefault('描述',[]).append(txt)
                elif cur:
                    lst=sk['fields'].setdefault(cur,[])
                    if lst and lst[-1]=='': lst[-1]=txt
                    else: lst.append(txt)
                for mm in re.finditer(r'(?:^|[。；！\s])([\u4e00-\u9fa5A-Za-z]{1,6})[：:]',txt):
                    lab=mm.group(1)
                    if ALIAS.get(lab,lab) not in ALL: st['unknown'][lab]+=1
        dtext=' '.join(sk['fields'].get('描述',[]))
        sk['upgrades']=len(re.findall(r'等级(?:到达|达到)\s*\d+\s*级时',dtext))
        st['skills'].append(sk)
    st['unknown']=dict(st['unknown'])
    return st

def main():
    import argparse
    ap = argparse.ArgumentParser(description='统计职业 docx 的技能块与字段覆盖')
    ap.add_argument('--pattern', default='基础职业-*.docx')
    ap.add_argument('--out', default='_scratch/docx_survey.json')
    args = ap.parse_args()
    data = [survey(p) for p in sorted(glob.glob(args.pattern))]
    os.makedirs(os.path.dirname(args.out) or '.', exist_ok=True)
    json.dump(data, open(args.out, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    total = sum(len(s['skills']) for s in data)
    cov = collections.Counter(); unk = collections.Counter(); colors = collections.Counter()
    multi = 0; markskill = 0; ups = 0
    for s in data:
        unk.update(s['unknown'])
        for sk in s['skills']:
            ups += sk['upgrades']
            for k, v in sk['fields'].items():
                if ' '.join(x for x in v if x).strip() not in ('', '-', '—'):
                    cov[k] += 1
            if sk['dots'] > 0: markskill += 1
            if len(set(sk['colors'])) > 1: multi += 1
            for c in sk['colors']: colors[c] += 1
    print('文档', len(data), '技能', total, '升级行', ups, '有标识', markskill, '多色', multi)
    for k in ALL:
        if cov[k]: print(f'  {k:6s} {cov[k]:5d} {cov[k]*100//max(total,1):3d}%')
    if unk: print('未识别标签:', dict(unk.most_common(20)))
    if colors: print('颜色:', ', '.join(f'{k}:{v}' for k, v in colors.most_common()))

if __name__ == '__main__':
    main()


