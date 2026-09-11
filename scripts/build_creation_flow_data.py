# -*- coding: utf-8 -*-
"""Extract 冒险者角色创建流程.xlsx into a web-friendly JS data file."""
import json
import os
import openpyxl
from openpyxl.utils import range_boundaries

SRC = '冒险者角色创建流程.xlsx'
OUT_DIR = os.path.join('斯诺德跑团', '资料库')
OUT = os.path.join(OUT_DIR, 'creation_flow_data.js')

def cell_text(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)

wb = openpyxl.load_workbook(SRC, data_only=True)
sheets = []
for ws in wb.worksheets:
    max_r = 0
    max_c = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value not in (None, ''):
                max_r = max(max_r, c.row)
                max_c = max(max_c, c.column)
    for rng in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(rng))
        max_r = max(max_r, r2)
        max_c = max(max_c, c2)
    if max_r == 0 or max_c == 0:
        continue
    rows = []
    for r in range(1, max_r + 1):
        vals = []
        for c in range(1, max_c + 1):
            vals.append(cell_text(ws.cell(r, c).value))
        rows.append(vals)
    merges = []
    for rng in ws.merged_cells.ranges:
        c1, r1, c2, r2 = range_boundaries(str(rng))
        if c2 <= max_c and r2 <= max_r:
            merges.append([c1, r1, c2, r2])
    sheets.append({'name': ws.title, 'maxCol': max_c, 'maxRow': max_r, 'rows': rows, 'merges': merges})

os.makedirs(OUT_DIR, exist_ok=True)
with open(OUT, 'w', encoding='utf-8') as f:
    f.write('// Auto-generated from 冒险者角色创建流程.xlsx\n')
    f.write('var CREATION_FLOW_DATA = ')
    json.dump({'sheets': sheets}, f, ensure_ascii=False, separators=(',', ':'))
    f.write(';\n')
print('written', OUT, 'sheets', len(sheets))
