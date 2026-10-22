#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从《个性与背景创建规则.xlsx》提取个性四类（特点/理念/羁绊/缺陷）的文字与色彩标识，
写入 职业页/数据/bg_personality.json 与 bg_personality_data.js 的 colors 字段。

用法:
  python -X utf8 scripts/build_bg_colors.py --report    # 只报告差异
  python -X utf8 scripts/build_bg_colors.py             # 写入
  python -X utf8 scripts/build_bg_colors.py --check     # 校验（供门禁）
"""
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _src_path():
    for i, a in enumerate(sys.argv):
        if a == "--src" and i + 1 < len(sys.argv):
            p = sys.argv[i + 1]
            return p if os.path.isabs(p) else os.path.join(ROOT, p)
    return os.path.join(ROOT, "个性与背景创建规则.xlsx")
OUT_JSON = os.path.join(ROOT, "职业页", "数据", "bg_personality.json")
OUT_JS = os.path.join(ROOT, "职业页", "数据", "bg_personality_data.js")

SKIP_SHEETS = {"支线剧情"}
FIELD_OF = {"特点": "traits", "理念": "ideals", "羁绊": "bonds", "缺陷": "flaws"}
FIELDS = ("traits", "ideals", "bonds", "flaws")
NONE_HEX = "#FFFFFF"  # 无填充 / 被清空的格子一律按白色

# ---- theme / fill 解析 -------------------------------------------------
try:
    from xml.etree import ElementTree as ET
    import colorsys
    _NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    _THEME_ORDER = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                    "accent4", "accent5", "accent6", "hlink", "folHlink"]
except Exception:  # pragma: no cover
    ET = None


def _apply_tint(hexv, tint):
    r = int(hexv[0:2], 16) / 255.0
    g = int(hexv[2:4], 16) / 255.0
    b = int(hexv[4:6], 16) / 255.0
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1.0 + tint) if tint < 0 else l * (1.0 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, max(0.0, min(1.0, l)), s)
    return "%02X%02X%02X" % (round(r * 255), round(g * 255), round(b * 255))


def theme_palette(wb):
    pal = {}
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return pal
    root = ET.fromstring(raw)
    scheme = root.find(".//" + _NS + "clrScheme")
    if scheme is None:
        return pal
    for idx, name in enumerate(_THEME_ORDER):
        node = scheme.find(_NS + name)
        if node is None:
            continue
        val = None
        for child in list(node):
            tag = child.tag.split("}")[-1]
            if tag == "srgbClr":
                val = child.get("val")
            elif tag == "sysClr":
                val = child.get("lastClr")
            if val:
                break
        if val:
            pal[str(idx)] = val.upper()
    return pal


def cell_hex(cell, wb, pal):
    sid = cell._style
    fid = sid.fillId if sid is not None else 0
    if fid is None or fid < 0 or fid >= len(wb._fills):
        return NONE_HEX
    fill = wb._fills[fid]
    info = vars(fill)
    if (info.get("patternType") or None) != "solid":
        return NONE_HEX
    fg = vars(info.get("fgColor") or {})
    ftype = fg.get("type")
    tint = fg.get("tint") or 0.0
    if ftype == "rgb":
        v = fg.get("rgb")
        v = v[-6:] if isinstance(v, str) and len(v) >= 6 else None
    elif ftype == "theme":
        tv = pal.get(str(fg.get("theme")))
        v = _apply_tint(tv, tint) if tv else None
    elif ftype == "indexed":
        try:
            from openpyxl.styles.colors import COLOR_INDEX
            v = COLOR_INDEX[int(fg.get("indexed"))]
        except Exception:
            v = None
    else:
        v = None
    if not v:
        return NONE_HEX
    return "#" + v.upper()


# ---- 表结构解析 --------------------------------------------------------
def parse_sheet(ws, wb, pal):
    """返回 {field: [(text, hex), ...]} 与特殊区块统计。"""
    out = {f: [] for f in FIELDS}
    specials = []
    block = None          # {cat, field, rows}
    for row in ws.iter_rows():
        b = row[1].value if len(row) > 1 else None
        c = row[2].value if len(row) > 2 else None
        d = row[3].value if len(row) > 3 else None
        bs = b.strip() if isinstance(b, str) else ""
        if isinstance(c, str) and c.strip() == "序列":
            block = {"cat": bs, "rows": []}
            specials.append(block)
            continue
        if block is None:
            continue
        if isinstance(c, (int, float)) and float(c) == int(c) and isinstance(d, str) and d.strip():
            if bs and not block["rows"]:
                block["cat"] = bs
            hexv = cell_hex(ws.cell(row=row[2].row, column=3), wb, pal)
            block["rows"].append((d.strip(), hexv))
    result = {f: [] for f in FIELDS}
    special_count = 0
    for blk in specials:
        field = FIELD_OF.get(blk["cat"])
        if field:
            result[field].extend(blk["rows"])
        elif blk["rows"]:
            special_count += len(blk["rows"])
    return result, special_count


def load_json():
    with open(OUT_JSON, encoding="utf-8") as f:
        return json.load(f)


def main():
    report = "--report" in sys.argv
    check = "--check" in sys.argv

    src = _src_path()
    print("== 数据源 ==")
    print("  " + os.path.relpath(src, ROOT) + "  (" + str(os.path.getsize(src)) + " B)")
    print()
    wb = openpyxl.load_workbook(src, data_only=True)
    pal = theme_palette(wb)
    sheets = [ws for ws in wb.worksheets if ws.title not in SKIP_SHEETS]
    data = load_json()

    problems = []
    text_diffs = []
    stat = {f: 0 for f in FIELDS}
    special_total = 0

    for ws in sheets:
        name = ws.title
        parsed, special_count = parse_sheet(ws, wb, pal)
        special_total += special_count
        if name not in data:
            problems.append("xlsx 表 %s 在 JSON 中不存在" % name)
            continue
        item = data[name]
        for field in FIELDS:
            rows = parsed[field]
            stat[field] += len(rows)
            texts = [t for t, _ in rows]
            cur = item.get(field) or []
            if texts != list(cur):
                text_diffs.append((name, field, list(cur), texts))
            if not check:
                item[field] = texts
                colors = item.setdefault("colors", {})
                colors[field] = [h for _, h in rows]
        if not check:
            item.setdefault("colors", {})

    # JSON 里有但 xlsx 里没有的职业
    xlsx_names = {ws.title for ws in sheets}
    for name in data:
        if name not in xlsx_names:
            problems.append("JSON 中 %s 在 xlsx 中不存在" % name)

    print("== 提取统计 ==")
    for f in FIELDS:
        print("  %-7s %4d 项" % (f, stat[f]))
    print("  合计      %4d 项（特殊选择区块 %d 项，不取色）" % (sum(stat.values()), special_total))
    print()
    if text_diffs:
        print("== 文本差异（%d 处）==" % len(text_diffs))
        for name, field, old, new in text_diffs:
            print("  [%s / %s] 旧 %d 项 → 新 %d 项" % (name, field, len(old), len(new)))
            import difflib
            for line in difflib.unified_diff(old, [t for t in new], lineterm="", n=0):
                if line.startswith(("---", "+++", "@@")):
                    continue
                print("    " + line)
    else:
        print("== 文本与现有 JSON 完全一致 ==")
    if problems:
        print()
        print("== 结构问题（%d）==" % len(problems))
        for p in problems:
            print("  " + p)

    # 颜色合法性
    print()
    from collections import Counter
    hist = Counter()
    missing = []
    for ws in sheets:
        name = ws.title
        parsed, _ = parse_sheet(ws, wb, pal)
        for field in FIELDS:
            for i, (txt, hexv) in enumerate(parsed[field]):
                hist[hexv] += 1
                if hexv == NONE_HEX and i >= 0:
                    pass
    print("== 色值分布 ==")
    for hexv, n in hist.most_common():
        print("  %-8s %4d" % (hexv, n))

    if check:
        bad = bool(text_diffs) or bool(problems)
        # 校验 JSON 里的 colors 是否与 xlsx 一致
        for ws in sheets:
            name = ws.title
            if name not in data:
                continue
            parsed, _ = parse_sheet(ws, wb, pal)
            colors = (data[name].get("colors") or {})
            for field in FIELDS:
                if colors.get(field) != [h for _, h in parsed[field]]:
                    bad = True
                    print("  FAIL 颜色不一致: %s / %s" % (name, field))
        if bad:
            print()
            print("FAIL: 个性背景颜色标识与 xlsx 不一致，请运行 python -X utf8 scripts/build_bg_colors.py")
            return 1
        print()
        print("OK: 个性背景颜色标识与 xlsx 一致（%d 项）" % sum(stat.values()))
        return 0

    if report or "--dry" in sys.argv:
        print()
        print("（--report 模式：未写入任何文件）")
        return 0

    with open(OUT_JSON, "w", encoding="utf-8", newline="\n") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")
    with open(OUT_JS, "w", encoding="utf-8", newline="\n") as f:
        f.write("var BG_PERSONALITY=" + json.dumps(data, ensure_ascii=False, indent=2) + ";\n")
    print()
    print("已写入 %s" % os.path.relpath(OUT_JSON, ROOT))
    print("已写入 %s" % os.path.relpath(OUT_JS, ROOT))
    return 0


if __name__ == "__main__":
    sys.exit(main())
