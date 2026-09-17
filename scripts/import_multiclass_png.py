#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把《兼职条件.png》写入 冒险者基础规则.xlsx 的「兼职规则」工作表（18 职业版）。

数据来源与可信度：
- 兼容矩阵：PNG 逐格像素解析（格内写列职业名=可兼职，'-'=不可兼职）；网格线自动检测
  （竖向 23 条 → 22 列 = 职业/属性/熟练度/其他 + 18 职业列；横向 20 条 → 表头 + 18 行）。
- 新增两行（召唤师/战舞者）的属性/熟练度/其他要求为人工转录（常量 NEW_ROWS），便于核对。
- 旧 16 行文本沿用表中既有值（本轮图与表逐行抽查一致），兼容矩阵整体以图为准覆盖。

用法：
  python scripts/import_multiclass_png.py --check   # 只解析 PNG 并与当前 xlsx 对比
  python scripts/import_multiclass_png.py --write   # 写入 xlsx（Excel COM，保留样式）
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PNG = ROOT / '兼职条件.png'
XLSX = ROOT / '冒险者基础规则.xlsx'
SHEET = '兼职规则'

CLASSES = ['蛮斗士', '战士', '法师', '猎人', '牧师', '圣骑士', '游荡者', '德鲁伊', '萨满祭司',
           '术士', '武僧', '吟游诗人', '魔契师', '奇械师', '守望者', '谋士', '召唤师', '战舞者']

# 新增行（人工转录自 PNG；战舞者的「舞蹈」在应用侧按 表演-舞蹈 归一，见 ref_subclass_reqs.py）
NEW_ROWS = {
    '召唤师': {
        'attr': '智力属性13，幸运属性14',
        'prof': '拥有奥秘、神秘学和机遇的熟练度共计+4',
        'other': '进行过系统性的咒法学派、降灵学科以及召唤学或契约学知识学习',
    },
    '战舞者': {
        'attr': '敏捷属性13，魅力属性13',
        'prof': '拥有体操、洞悉和舞蹈的熟练度共计+4',
        'other': '-',
    },
}
PLANNED = {'战舞者'}  # 未开放职业：表内保留，帮助页标注「未开放」，兼职候选不出现


def detect_grid(png: Path):
    """检测表格网格线，返回 (col_lines, row_lines)。"""
    import numpy as np
    from PIL import Image
    a = np.array(Image.open(png).convert('RGB')).astype(int)
    h, w, _ = a.shape
    dark = (a[:, :, 0] < 200) & (a[:, :, 1] < 170) & (a[:, :, 2] < 140)
    col_ink = dark.sum(axis=0)
    row_ink = dark.sum(axis=1)

    def group(idx):
        out = []
        for x in idx:
            if out and x - out[-1][-1] <= 3:
                out[-1].append(x)
            else:
                out.append([x])
        return [int(sum(g) / len(g)) for g in out]

    cols = group([x for x in range(w) if col_ink[x] > h * 0.6])
    rows = group([y for y in range(h) if row_ink[y] > w * 0.6])
    return cols, rows


def extract_matrix(png: Path):
    import numpy as np
    from PIL import Image
    cols, rows = detect_grid(png)
    if len(cols) != 4 + len(CLASSES) + 1:
        raise SystemExit('网格检测异常：竖向线 %d 条（期望 %d）：%s' % (len(cols), 4 + len(CLASSES) + 1, cols))
    if len(rows) != len(CLASSES) + 2:
        raise SystemExit('网格检测异常：横向线 %d 条（期望 %d）：%s' % (len(rows), len(CLASSES) + 2, rows))
    a = np.array(Image.open(png).convert('RGB')).astype(int)
    ink = (a[:, :, 0] < 235) & (a[:, :, 1] < 215) & (a[:, :, 2] < 190)
    mat: dict[str, dict[str, bool]] = {}
    for ri, rname in enumerate(CLASSES):
        y0, y1 = rows[ri + 1], rows[ri + 2]
        mat[rname] = {}
        for ci, cname in enumerate(CLASSES):
            x0, x1 = cols[4 + ci], cols[5 + ci]
            cnt = int(ink[y0 + 4:y1 - 4, x0 + 4:x1 - 4].sum())
            mat[rname][cname] = cnt > 40
    bad = [c for c in CLASSES if not mat[c][c]]
    if bad:
        raise SystemExit('对角线自兼容断言失败：%s' % bad)
    return mat, cols, rows


def read_xlsx_matrix() -> dict:
    import openpyxl
    ws = openpyxl.load_workbook(XLSX, data_only=True)[SHEET]
    out = {}
    for r in range(3, ws.max_row + 1):
        name = str(ws.cell(r, 2).value or '').strip()
        if not name:
            continue
        row = {}
        for ci, cn in enumerate(CLASSES):
            row[cn] = str(ws.cell(r, 6 + ci).value or '').strip() == cn
        out[name] = row
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--write', action='store_true')
    ap.add_argument('--png', default=str(DEFAULT_PNG))
    args = ap.parse_args()

    png = Path(args.png)
    if not png.exists():
        raise SystemExit('未找到 %s' % png)
    mat, cols, rows = extract_matrix(png)
    print('网格: 竖 %d 条 / 横 %d 条 → %d×%d 矩阵' % (len(cols), len(rows), len(CLASSES), len(CLASSES)))
    asym = [(r, c) for r in CLASSES for c in CLASSES if r != c and mat[r][c] != mat[c][r]]
    print('非对称关系 %d 对（按作者原图保留）：%s'
          % (len(asym), '、'.join('%s×%s' % p for p in asym[:8]) or '无'))
    for r in CLASSES:
        print('  %-5s%s 可兼职: %s' % (r, '（未开放）' if r in PLANNED else '      ',
                                     '、'.join(c for c in CLASSES if mat[r][c]) or '（无）'))

    cur = read_xlsx_matrix()
    diffs = []
    for r, row in mat.items():
        for c, v in row.items():
            if r in cur and c in cur.get(r, {}):
                if cur[r].get(c) != v:
                    diffs.append((r, c, cur[r].get(c), v))
    print('与当前 xlsx 的矩阵差异: %d 处%s' % (len(diffs), '（旧 16×16 部分完全一致）' if not diffs else ''))
    for r, c, old, new in diffs[:20]:
        print('   %s × %s: %s → %s' % (r, c, old, new))
    missing_rows = [c for c in CLASSES if c not in cur]
    if missing_rows:
        print('xlsx 待新增职业行:', '、'.join(missing_rows))
    if args.check or not args.write:
        print('（未写入；需要写入请加 --write）')
        return 0

    try:
        import win32com.client as win32  # type: ignore
    except ImportError:
        raise SystemExit('需要 pywin32 才能保留 Excel 样式；请安装后再运行')

    excel = win32.DispatchEx('Excel.Application')
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(XLSX))
        try:
            ws = wb.Worksheets(SHEET)
            base_row, base_col = 18, 21          # 现有 16 职业：行 3..18、列 6..21（F..U）
            for i in range(len(CLASSES) - 16):
                ws.Rows(base_row).Copy()
                ws.Rows(base_row + 1 + i).PasteSpecial(-4122)   # xlPasteFormats
            for i in range(len(CLASSES) - 16):
                ws.Columns(base_col).Copy()
                ws.Columns(base_col + 1 + i).PasteSpecial(-4122)
            excel.CutCopyMode = False
            last_col = 5 + len(CLASSES)              # 18 职业 → 第 23 列（W）
            last_addr = get_column_letter(last_col)   # COM 的 Address 在此为属性，改用 openpyxl 工具函数
            try:
                ws.Range('F2:U2').UnMerge()
            except Exception:
                pass
            ws.Range('F2:%s2' % last_addr).Merge()
            ws.Range('F2').Value = '兼容要求'
            ws.Range('F2').Copy()
            ws.Range('F2:%s2' % last_addr).PasteSpecial(-4122)
            excel.CutCopyMode = False
            for ri, rname in enumerate(CLASSES):
                row = 3 + ri
                if rname in NEW_ROWS:
                    ws.Cells(row, 2).Value = rname
                    ws.Cells(row, 3).Value = NEW_ROWS[rname]['attr']
                    ws.Cells(row, 4).Value = NEW_ROWS[rname]['prof']
                    ws.Cells(row, 5).Value = NEW_ROWS[rname]['other']
                for ci, cname in enumerate(CLASSES):
                    ws.Cells(row, 6 + ci).Value = cname if mat[rname][cname] else '-'
            wb.Save()
        finally:
            wb.Close(SaveChanges=False)
    finally:
        excel.Quit()
    print('✅ 已写入 %s（%d 职业，矩阵 %d×%d）' % (XLSX.name, len(CLASSES), len(CLASSES), len(CLASSES)))
    return 0


if __name__ == '__main__':
    sys.exit(main())
