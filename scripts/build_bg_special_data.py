# -*- coding: utf-8 -*-
"""把《个性与背景创建规则.xlsx》里此前未入库的内容补进背景数据。

补录内容（v1.0.7269）：
- 侍僧「神祇简介」的 11 条神祇描述
- 骗子「偏好骗局」的 6 条对应装备
- 恶棍「罪名」6 条罪名 + 接头人
- 乐师 / 艺人「名望加成」表（3 档）
- 驯兽师「名望加成」里的 8 种动物伙伴
- 隐士「隐居原因」8 条
- 士兵「专职」8 条（熟练度加成 + 额外装备）
- 外乡人「原因」6 条（身份 + 介绍）
- 教授「学术领域」6 条（熟练度加成 + 额外装备）

输出：职业页/数据/bg_personality.json 与 bg_personality_data.js（两者保持同步）
用法：python scripts/build_bg_special_data.py [--check]
"""
import io
import json
import os
import subprocess
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, '个性与背景创建规则.xlsx')
OUT_JSON = os.path.join(ROOT, '职业页', '数据', 'bg_personality.json')
OUT_JS = os.path.join(ROOT, '职业页', '数据', 'bg_personality_data.js')

ORDER = ['name', 'traits', 'ideals', 'bonds', 'flaws', 'equipment', 'deities',
         'contacts', 'scam_types', 'scam_details', 'mission_channels',
         'academic_domains', 'academic_details', 'deity_details',
         'crimes', 'crime_details', 'fame_tiers', 'companions',
         'seclusion_reasons', 'military_roles', 'foreign_origins',
         'desc', 'hp', 'base_profs', 'special_profs', 'other', 'gold', 'trait_desc']

# 各背景的特殊段落：标签 -> 目标字段名
SECTION_FIELDS = {
    '侍僧': [('神祇简介', 'deity_details')],
    '骗子': [('偏好骗局', 'scam_details')],
    '恶棍': [('罪名', 'crime_details')],
    '乐师': [('名望加成', 'fame_tiers')],
    '艺人': [('名望加成', 'fame_tiers')],
    '驯兽师': [('名望加成', 'companions')],
    '隐士': [('隐居原因', 'seclusion_reasons')],
    '士兵': [('专职', 'military_roles')],
    '外乡人': [('原因', 'foreign_origins')],
    '教授': [('学术领域', 'academic_details')],
}

# 段落内每列的含义（列号 -> 字段名）；键为“目标字段名”
COLUMN_MEANING = {
    'deity_details': {3: 'name', 5: 'desc'},
    'scam_details': {4: 'name', 5: 'gear'},
    'crime_details': {4: 'name', 5: 'contact'},
    'fame_tiers': {3: 'level', 4: 'bonus'},
    'companions': {4: 'text'},
    'seclusion_reasons': {4: 'text'},
    'military_roles': {4: 'name', 5: 'prof', 6: 'gear'},
    'foreign_origins': {4: 'name', 5: 'desc'},
    'academic_details': {4: 'name', 5: 'prof', 6: 'gear'},
}


def text(v):
    if v is None:
        return ''
    return str(v).strip()


# 段落表头词（跳过表头行；侍僧「神祇简介」的首条与段落标题同处一行，需要保留）
HEADER_WORDS = {'序列', '编号', '名称', '骗局', '装备', '接头人', '主要联系人', '委托人',
                '等级', '增益', '种类', '原因', '身份', '介绍', '学科', '职务',
                '熟练度加成', '额外装备', '领域', '效果', '罪名', '魅力相关检定增益'}


def load_current():
    """用 node 读取现有 JS 数据（JSON 文件可能缺少后加字段）。"""
    js = subprocess.run(['node', '-e',
                         "const fs=require('fs');const d=new Function(fs.readFileSync(process.argv[1],'utf8')+'; return BG_PERSONALITY;')();process.stdout.write(JSON.stringify(d));",
                         OUT_JS], check=True, capture_output=True, text=True, encoding='utf-8')
    return json.loads(js.stdout)


def extract_sections():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    result = {}
    for sheet, specs in SECTION_FIELDS.items():
        ws = wb[sheet]
        for label, field in specs:
            start = None
            for r in range(1, ws.max_row + 1):
                if text(ws.cell(row=r, column=2).value).replace('\n', '') == label:
                    start = r
                    break
            if start is None:
                continue
            meaning = COLUMN_MEANING[field]
            rows = []
            for r in range(start, ws.max_row + 1):
                head = text(ws.cell(row=r, column=2).value).replace('\n', '')
                if r > start and head in ('个性', '特点', '理念', '羁绊', '缺陷'):
                    break
                raw = {}
                for c, key in meaning.items():
                    v = text(ws.cell(row=r, column=c).value)
                    if v:
                        raw[key] = v
                if raw and all(v in HEADER_WORDS for v in raw.values()):
                    continue  # 表头行（部分段落的标题行本身就是表头行）
                rec = {k: v for k, v in raw.items() if v not in HEADER_WORDS}
                if rec:
                    rows.append(rec)
            if not rows:
                continue
            if field in ('companions', 'seclusion_reasons', 'fame_tiers'):
                if field == 'fame_tiers':
                    result.setdefault(sheet, {})[field] = [{'level': x.get('level', ''), 'bonus': x.get('bonus', '')} for x in rows]
                else:
                    result.setdefault(sheet, {})[field] = [x['text'] for x in rows]
            else:
                result.setdefault(sheet, {})[field] = rows
                extra = {'crime_details': ('crimes', 'name'), 'military_roles': ('roles', 'name'),
                         'foreign_origins': ('origins', 'name')}
                if field in extra:
                    key, src_key = extra[field]
                    result[sheet].setdefault(key, [x[src_key] for x in rows if src_key in x])
    return result


def ordered(item):
    out = {}
    for k in ORDER:
        if k in item:
            out[k] = item[k]
    for k, v in item.items():
        if k not in out:
            out[k] = v
    return out


def main():
    check = '--check' in sys.argv
    base = load_current()
    specials = extract_sections()
    for name, fields in specials.items():
        base.setdefault(name, {})
        base[name].update(fields)
    data = {k: ordered(v) for k, v in base.items()}
    js = 'var BG_PERSONALITY=' + json.dumps(data, ensure_ascii=False, indent=2) + ';\n'
    jsn = json.dumps(data, ensure_ascii=False, indent=2) + '\n'
    if check:
        cur_js = io.open(OUT_JS, encoding='utf-8').read()
        cur_json = io.open(OUT_JSON, encoding='utf-8').read()
        ok = (cur_js == js and cur_json == jsn)
        print('OK: 背景数据与源文件一致' if ok else 'FAIL: 背景数据与源文件不一致')
        return 0 if ok else 1
    io.open(OUT_JS, 'w', encoding='utf-8', newline='\n').write(js)
    io.open(OUT_JSON, 'w', encoding='utf-8', newline='\n').write(jsn)
    total = sum(len(v) for fields in specials.values() for v in fields.values())
    print('backgrounds updated:', len(specials))
    for name, fields in specials.items():
        print('  %s: %s' % (name, ', '.join('%s(%d)' % (k, len(v)) for k, v in fields.items())))
    print('written:', OUT_JSON, '|', OUT_JS)
    return 0


if __name__ == '__main__':
    sys.exit(main())
