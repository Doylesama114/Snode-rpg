#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把谋士写入角色面板数据（panel_data.js）：REF_CLASSES / CLASS_WEAPON_PROF_DOCX / CLASS_WEAPON_PROFS。

数据来源：
- 数值/熟练项：职业页/数据/classes_data.js（由谋士 docx 转录）
- 专长（6 战斗风格）与起始特性（4 选 2）：职业页/数据/谋士.json + docx 风格简介
用法：
  python scripts/apply_strategist_panel_entries.py --check
  python scripts/apply_strategist_panel_entries.py --write
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

ROOT = Path(__file__).resolve().parent.parent
PANEL = ROOT / '斯诺德跑团' / 'panel_data.js'
PANEL_MIRROR = ROOT / 'electron-app' / '斯诺德跑团' / 'panel_data.js'
CLASSES_DATA = ROOT / '职业页' / '数据' / 'classes_data.js'
MOUSHI_JSON = ROOT / '职业页' / '数据' / '谋士.json'
MOUSHI_DOCX = ROOT / '基础职业-谋士.docx'

CLASS = '谋士'
WEAPON_DOCX = '匕首、手弩、简易武器'
# 匕首→剑类、手弩→弓箭、简易武器→简易（与术士同款武器表）
WEAPON_CATS = ['剑类', '弓箭', '简易']
STYLE_DESCS = {
    '权谋': '洞察人心、操弄信息、统御他人来改变局势',
    '军团': '战场上的指挥才能，通过调度、布阵、增益与战术配合放大团队战力',
    '先见': '对战局的前瞻性洞察与精密预判，落子前已推演数步',
    '鸩毒': '对毒物与毒计的极致掌握，让敌人在不知不觉中衰弱',
    '混乱': '以干扰、误导与制造无序为核心，在敌阵内部播撒猜忌与恐慌',
    '博物': '对世间万物的广博认知与精密分析，以知识本身为武器',
}


def js_const(path: Path, name: str, kind: str = 'var'):
    text = path.read_text(encoding='utf-8')
    m = re.search(r'%s\s+%s\s*=\s*(\{.*?\});' % (re.escape(kind), re.escape(name)), text, re.S)
    if not m:
        raise SystemExit('未找到 %s in %s' % (name, path))
    return json.loads(m.group(1)), m.group(1)


def load_classes_data() -> dict:
    text = CLASSES_DATA.read_text(encoding='utf-8')
    arr = json.loads(text[text.index('['):text.rindex(']') + 1])
    return next(c for c in arr if c['name'] == CLASS)


def style_desc(style: str) -> str:
    """优先取 docx 风格简介首句。"""
    try:
        import docx
        doc = docx.Document(str(MOUSHI_DOCX))
        for i, p in enumerate(doc.paragraphs):
            t = ' '.join(p.text.split())
            if t == f'{style}风格':
                for q in doc.paragraphs[i + 1:i + 4]:
                    d = ' '.join(q.text.split())
                    if d and not d.startswith('---'):
                        first = d.split('——')[0]
                        if '。' in first:
                            first = first.split('。')[0] + '。'
                        return first
    except Exception:
        pass
    return STYLE_DESCS.get(style, '')


def build_entry() -> dict:
    cd = load_classes_data()
    moushi = json.loads(MOUSHI_JSON.read_text(encoding='utf-8'))
    styles = []
    seen = []
    for s in moushi['skills']:
        st = s.get('style')
        if st and st not in seen:
            seen.append(st)
    for st in seen:
        styles.append({'name': st, 'desc': style_desc(st)})
    short = {
        '交友术': '使一名角色将你视作好友',
        '战术部署': '为友方角色分配战术增益效果',
        '毒刃': '在武器上淬毒并发起一击',
        '离间': '离间敌人和其盟友，令关系暂时恶化',
    }
    features = []
    for s in moushi['skills']:
        if s.get('type') == 'starting':
            desc = short.get(s['name']) or (s.get('fields', {}).get('描述') or '').strip()
            features.append({'name': s['name'], 'desc': desc[:40]})
    return {
        'id': CLASS,
        'name': CLASS,
        'key_attr': cd['关键属性'],
        'armor': cd['护甲'],
        'weapons': cd['武器'],
        'saves': cd['豁免'].split('、'),
        'skills': cd['技巧'],
        'specializations': styles,
        'starting_features': features,
        'starting_choice': 2,
        'hp_formula': cd['hp_formula'],
        'fp_formula': cd['fp_formula'],
    }


def read_ref(text: str) -> dict:
    m = re.search(r"const REF_CLASSES = JSON\.parse\('(.*?)'\);", text, re.S)
    if not m:
        raise SystemExit('未找到 REF_CLASSES')
    return json.loads(m.group(1).replace("\\'", "'"))


def write_ref(text: str, ref: dict) -> str:
    body = json.dumps(ref, ensure_ascii=False).replace("'", "\\'")
    return re.sub(r"const REF_CLASSES = JSON\.parse\('.*?'\);",
                  "const REF_CLASSES = JSON.parse('%s');" % body, text, count=1, flags=re.S)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--write', action='store_true')
    args = ap.parse_args()

    entry = build_entry()
    text = PANEL.read_text(encoding='utf-8')
    ref = read_ref(text)
    prof_docx, _ = js_const(PANEL, 'CLASS_WEAPON_PROF_DOCX')
    prof_cat, _ = js_const(PANEL, 'CLASS_WEAPON_PROFS')

    has_ref = CLASS in ref
    has_docx = CLASS in prof_docx
    has_cat = CLASS in prof_cat
    print('REF_CLASSES: %s' % ('已写入' if has_ref else '缺失'))
    print('CLASS_WEAPON_PROF_DOCX: %s' % ('已写入' if has_docx else '缺失'))
    print('CLASS_WEAPON_PROFS: %s' % ('已写入' if has_cat else '缺失'))
    print('专长 %d 条：%s' % (len(entry['specializations']), '、'.join(x['name'] for x in entry['specializations'])))
    print('起始特性 %d 条：%s' % (len(entry['starting_features']), '、'.join(x['name'] for x in entry['starting_features'])))

    if args.check:
        ok = has_ref and has_docx and has_cat
        if ok:
            if ref[CLASS].get('hp_formula') != entry['hp_formula'] or ref[CLASS].get('weapons') != WEAPON_DOCX:
                ok = False
                print('FAIL: REF_CLASSES 内容与预期不一致')
        print('OK' if ok else 'FAIL')
        return 0 if ok else 1
    if not args.write:
        ap.print_help()
        return 0

    ref[CLASS] = entry
    text = write_ref(text, ref)
    prof_docx[CLASS] = WEAPON_DOCX
    text = re.sub(r'var CLASS_WEAPON_PROF_DOCX=\{[^;]+\};',
                  'var CLASS_WEAPON_PROF_DOCX=%s;' % json.dumps(prof_docx, ensure_ascii=False),
                  text, count=1)
    prof_cat[CLASS] = WEAPON_CATS
    text = re.sub(r'var CLASS_WEAPON_PROFS=\{[^;]+\};',
                  'var CLASS_WEAPON_PROFS=%s;' % json.dumps(prof_cat, ensure_ascii=False),
                  text, count=1)
    PANEL.write_text(text, encoding='utf-8', newline='')
    if PANEL_MIRROR.exists():
        shutil.copyfile(PANEL, PANEL_MIRROR)
    print('✅ 面板谋士条目已写入')
    return 0


if __name__ == '__main__':
    sys.exit(main())
