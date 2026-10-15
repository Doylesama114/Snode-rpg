# -*- coding: utf-8 -*-
"""从 斯诺德稀有装备列表.xlsx 生成 职业页/数据/equipment_catalog.json（保留 price 与结构/格式）"""
import openpyxl, json, sys
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path('.').resolve()
XLSX = ROOT / '斯诺德稀有装备列表.xlsx'
JSONP = ROOT / '职业页' / '数据' / 'equipment_catalog.json'
DATAP = ROOT / '职业页' / '数据' / 'equipment_catalog_data.js'
APPLY = '--write' in sys.argv
wb = openpyxl.load_workbook(XLSX, data_only=True)
ws = wb.worksheets[0]
def cell(r, c):
    v = ws.cell(row=r, column=c).value
    return '' if v is None else str(v).strip()
rows = []
for r in range(2, ws.max_row + 1):
    name = cell(r, 1)
    if not name:
        continue
    rows.append({'name': name, 'category': cell(r, 2), 'subcategory': cell(r, 3), 'type': cell(r, 4), 'weight': cell(r, 5), 'requirement': cell(r, 6), 'effect': cell(r, 7)})
print('xlsx 条目数: %d' % len(rows))
old = json.loads(JSONP.read_text(encoding='utf-8'))
old_index = {}
cat_order = list(old.keys())
for cat, items in old.items():
    for it in items:
        old_index[it['name']] = it
print('旧 JSON: %d 类别 / %d 条目' % (len(old), sum(len(v) for v in old.values())))
keep_price = {n: it.get('price', '不可购买') for n, it in old_index.items()}
new = {}
missing_price = []
for cat in cat_order:
    new[cat] = []
for row in rows:
    cat = row['category'] or '未分类'
    if cat not in new:
        new[cat] = []
    p = keep_price.get(row['name'])
    if p is None:
        missing_price.append(row['name'])
        p = '不可购买'
    new[cat].append({'name': row['name'], 'category': cat, 'subcategory': row['subcategory'], 'type': row['type'], 'weight': row['weight'], 'requirement': row['requirement'], 'effect': row['effect'], 'price': p})
new = {k: v for k, v in new.items() if v}
oldnames = set(old_index)
newnames = set(r['name'] for r in rows)
print('')
print('新增条目: %s' % (sorted(newnames - oldnames) or '无'))
print('被移除条目: %s' % (sorted(oldnames - newnames) or '无'))
print('继承不到 price 的: %s' % (missing_price or '无'))
print('新 JSON: %d 类别 / %d 条目' % (len(new), sum(len(v) for v in new.values())))
for k, v in new.items():
    print('    %s: %d' % (k, len(v)))
if APPLY:
    JSONP.write_text(json.dumps(new, ensure_ascii=False, indent=4) + chr(10), encoding='utf-8')
    DATAP.write_text('var EQUIPMENT_CATALOG = ' + json.dumps(new, ensure_ascii=False, indent=4) + ';' + chr(10), encoding='utf-8')
    print('')
    print('已写入: %s' % JSONP)
    print('已写入: %s' % DATAP)
else:
    print('')
    print('（预览模式，未写入；加 --write 执行）')