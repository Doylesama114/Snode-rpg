#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""谋士版本 gap report：把三份源文件（进阶途径 docx / 谋士 docx / 兼职 PNG→xlsx）
与当前应用数据逐项对照，产出实施清单与验收基线。

产出：
  scripts/_strategist_gap_report.json   结构化报告（机器可读）
用法：
  python scripts/build_strategist_gap_report.py
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import defaultdict, OrderedDict
from datetime import date
from pathlib import Path

import openpyxl  # noqa: E402

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

ROOT = Path(__file__).resolve().parent.parent
MOUSHI_DOCX = ROOT / '基础职业-谋士.docx'
ADV_DOCX = ROOT / '《基础职业进阶途径》.docx'
XLSX = ROOT / '冒险者基础规则.xlsx'
OUT = Path(__file__).resolve().parent / '_strategist_gap_report.json'

CLASSES = ['蛮斗士', '战士', '法师', '猎人', '牧师', '圣骑士', '游荡者', '德鲁伊', '萨满祭司',
           '术士', '武僧', '吟游诗人', '魔契师', '奇械师', '守望者', '谋士']
STYLES = ['权谋', '军团', '先见', '鸩毒', '混乱', '博物']
APP = ROOT / '斯诺德跑团'
PAGE = ROOT / '职业页'


def sha1(p: Path) -> str:
    h = hashlib.sha1()
    h.update(p.read_bytes())
    return h.hexdigest()


def rows_of(tb):
    out = []
    for row in tb.rows:
        cells, seen = [], set()
        for c in row.cells:
            if id(c._tc) in seen:
                continue
            seen.add(id(c._tc))
            cells.append(' '.join(c.text.split()))
        out.append(cells)
    return out


def parse_moushi() -> dict:
    """解析谋士 docx：数值/熟练项/装备/专长/起始特性/天赋树。"""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(MOUSHI_DOCX))
    out = {'file': str(MOUSHI_DOCX.name), 'sha1': sha1(MOUSHI_DOCX),
           'paragraphs': [], 'skills': [], 'tables': 0}
    style = tier = choice = None
    zone = 'meta'
    for child in d.element.body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            t = ' '.join(Paragraph(child, d).text.split())
            if not t:
                continue
            out['paragraphs'].append(t)
            if t == '谋士天赋树':
                zone, style, tier, choice = 'talents', None, None, None
            elif t == '起始特性':
                zone, style, tier, choice = 'starting', None, None, None
                out['_pending_start_list'] = True
            elif t == '初始专长':
                zone, style, tier, choice = 'features', None, None, None
            elif zone == 'talents' and t.endswith('风格') and len(t) <= 10:
                style, tier, choice = t[:-2], None, None
            elif zone == 'talents' and '天赋树' in t[:10]:
                tier, choice = t.split('天赋树')[0], None
            elif zone == 'talents' and t.startswith('抉择'):
                choice = t
            continue
        if tag != 'tbl':
            continue
        out['tables'] += 1
        rows = rows_of(Table(child, d))
        flat = ' '.join(' '.join(r) for r in rows)
        name = rows[0][0] if rows and rows[0] else ''
        is_skill = ('施展时间' in flat) or ('关键词' in flat and '持续时间' in flat)
        # 起始特性清单表（紧跟「起始特性」段落的第一张表）→ 记录特性名
        if zone == 'starting' and out.get('_pending_start_list') and len(rows) > 1:
            out['starting_list'] = [r[0] for r in rows if r and r[0]]
            out['_pending_start_list'] = False
        if is_skill and name and len(name) <= 20 and name not in ('关键词', '名称'):
            if zone == 'talents':
                out['skills'].append({'name': name, 'style': style or '', 'tier': tier,
                                      'choice': choice, 'rows': len(rows)})
            elif zone == 'starting':
                out['skills'].append({'name': name, 'style': '起始特性', 'tier': None,
                                      'choice': None, 'rows': len(rows)})
    # 职业专长（初始专长：名称 + 描述成对出现）
    ps = out['paragraphs']
    feats = []
    if '初始专长' in ps:
        i = ps.index('初始专长') + 1
        while i < len(ps) and ps[i] != '战斗风格':
            name = ps[i]
            desc = ps[i + 1] if i + 1 < len(ps) else ''
            if name and not name.startswith('---') and not name.startswith('谋士的关键属性'):
                feats.append({'name': name, 'desc': desc})
                i += 2
            else:
                i += 1
    out['features'] = feats
    out['starting_features'] = out.get('starting_list', [])
    out.pop('_pending_start_list', None)
    # 关键数值（段落抽取）
    def grab(prefix):
        for p in out['paragraphs']:
            if p.startswith(prefix):
                return p
        return ''
    out['summary'] = {
        'role': next((p for p in out['paragraphs'] if p.startswith('职责定位')), ''),
        'hp': [p for p in out['paragraphs'] if p.startswith('首级生命值') or p.startswith('升级生命值')],
        'fp': [p for p in out['paragraphs'] if p.startswith('首级疲劳值') or p.startswith('升级疲劳值')],
        'profs': [p for p in out['paragraphs'] if re.match(r'^(护甲|武器|豁免|技艺)[:：]', p)],
        'equip_sets': [p for p in out['paragraphs'] if re.match(r'^[A-D]·', p)],
        'features': [],
    }
    out['summary']['features'] = out.get('features', [])
    return out


def parse_advancements() -> dict:
    """解析进阶途径 docx：谋士 25 条 + 全表同名来源并集。"""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(ADV_DOCX))
    section = None
    cards = []
    for child in d.element.body.iterchildren():
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            t = ' '.join(Paragraph(child, d).text.split())
            if t.endswith('进阶途径') and len(t) <= 12:
                section = t[:-4]
            continue
        if tag != 'tbl':
            continue
        rows = rows_of(Table(child, d))
        if not rows or '属性值需求' not in ' '.join(rows[0]):
            continue
        hdr, name_r = rows[1] if len(rows) > 1 else [], rows[2] if len(rows) > 2 else []
        attrs = {}
        if len(name_r) >= 10 and len(hdr) >= 10:
            for idx in range(2, 10):
                v = name_r[idx]
                if v and v not in ('-',):
                    attrs[hdr[idx]] = v
        src = ''
        conds = []
        for r in rows[3:]:
            if not r:
                continue
            if r[0] == '来源' and len(r) > 1:
                src = r[1]
            elif r[0].isdigit() and len(r) > 1:
                conds.append(r[1])
        cards.append({'section': section, 'name': name_r[0] if name_r else '',
                      'attrs': attrs, 'source': src, 'conditions': conds})
    moushi = [c for c in cards if c['section'] == '谋士']
    byname = defaultdict(list)
    for c in cards:
        byname[c['name']].append(c)
    unified = {}
    for name, cs in byname.items():
        if all(c['section'] == '通用' for c in cs):
            continue
        srcs = []
        for c in cs:
            for t in c['source'].split('、'):
                t = t.strip()
                if t and t != '全职业' and t not in srcs:
                    srcs.append(t)
        srcs.sort(key=lambda n: CLASSES.index(n) if n in CLASSES else 99)
        if len({c['source'] for c in cs}) > 1:
            unified[name] = {'unified': srcs, 'variants': sorted({c['source'] for c in cs})}
    return {'file': str(ADV_DOCX.name), 'sha1': sha1(ADV_DOCX),
            'moushi_cards': [{'name': c['name'], 'attrs': c['attrs'], 'source': c['source'],
                              'conditions': len(c['conditions'])} for c in moushi],
            'moushi_new': [c['name'] for c in moushi if len(byname[c['name']]) == 1],
            'unified_sources': unified,
            'universal_count': sum(1 for c in cards if c['section'] == '通用')}


def parse_multiclass() -> dict:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb['兼职规则']
    data = {}
    for r in range(3, 19):
        name = ws.cell(r, 2).value
        if not name:
            continue
        name = str(name).strip()
        compat = []
        for c in range(6, 22):
            v = ws.cell(r, c).value
            if v and str(v).strip() not in ('-', '') and str(v).strip() in CLASSES:
                compat.append(str(v).strip())
        data[name] = {
            'attr': str(ws.cell(r, 3).value or '').strip(),
            'prof': str(ws.cell(r, 4).value or '').strip(),
            'other': str(ws.cell(r, 5).value or '').strip(),
            'compatible': compat,
            'incompatible': [c for c in CLASSES if c not in compat],
        }
    asym = [(a, b) for a in data for b in data
            if (b in data[a]['compatible']) != (a in data[b]['compatible'])]
    return {'file': str(XLSX.name), 'classes': len(data), 'data': data,
            'asymmetries': asym}


def parse_app_state() -> dict:
    st = {'classes_data': [], 'class_features': [], 'equipment_data': [],
          'skill_effects': [], 'page_json': [], 'ref_subclass_reqs': [], 'chargen_adv_paths': []}
    cd = PAGE / '数据' / 'classes_data.js'
    if cd.exists():
        st['classes_data'] = re.findall(r'"name":\s*"([^"]+)"', cd.read_text(encoding='utf-8'))
    cf = PAGE / '数据' / 'class_features.json'
    if cf.exists():
        st['class_features'] = list(json.loads(cf.read_text(encoding='utf-8')).get('classes', {}).keys())
    eq = PAGE / '数据' / 'equipment_data.js'
    if eq.exists():
        st['equipment_data'] = re.findall(r'"([^"]+)":\s*\[', eq.read_text(encoding='utf-8'))[:32]
    st['skill_effects'] = sorted(p.name for p in APP.glob('skill_effects_*.json'))
    st['page_json'] = sorted(p.stem for p in (PAGE / '数据').glob('*.json'))
    pd = APP / 'panel_data.js'
    if pd.exists():
        m = re.search(r'var REF_SUBCLASS_REQS = (\{.*?\});', pd.read_text(encoding='utf-8'), re.S)
        if m:
            try:
                st['ref_subclass_reqs'] = list(json.loads(m.group(1)).keys())
            except Exception:
                st['ref_subclass_reqs'] = re.findall(r'"([^"]+)":\{"attrs"', m.group(1))
    cap = APP / 'chargen_adv_paths.js'
    if cap.exists():
        m = re.search(r'window\.CHARGEN_ADV_PATHS = (\{.*\});', cap.read_text(encoding='utf-8'), re.S)
        if m:
            try:
                st['chargen_adv_paths'] = sorted(json.loads(m.group(1)).keys())
            except Exception:
                pass
    return st


def main() -> int:
    moushi = parse_moushi()
    adv = parse_advancements()
    mc = parse_multiclass()
    app = parse_app_state()

    skills = moushi['skills']
    starting_features = moushi.get('starting_features') or [s['name'] for s in skills if s['style'] == '起始特性']
    by_tier = defaultdict(list)
    for s in skills:
        if s['style'] == '起始特性':
            continue
        by_tier['%s/%s' % (s['style'], s['tier'])].append(s['name'])
    report = OrderedDict()
    report['generatedAt'] = date.today().isoformat()
    report['sources'] = {'moushi_docx': moushi['sha1'], 'adv_docx': adv['sha1'],
                         'multiclass_xlsx': sha1(XLSX)}
    report['multiclass'] = {
        'classes': mc['classes'],
        'moushi': mc['data'].get('谋士', {}),
        'asymmetries': mc['asymmetries'],
        'app_ref_subclass_reqs': app['ref_subclass_reqs'],
        'app_has_moushi': '谋士' in app['ref_subclass_reqs'],
    }
    report['advancements'] = {
        'moushi_cards': len(adv['moushi_cards']),
        'moushi_new': adv['moushi_new'],
        'unified_sources': adv['unified_sources'],
        'universal_count': adv['universal_count'],
        'app_page_json_has_moushi': '谋士' in app['page_json'],
        'app_chargen_adv_has_moushi': '谋士' in app['chargen_adv_paths'],
    }
    report['class'] = {
        'name': '谋士',
        'summary': moushi['summary'],
        'skill_count': len(skills),
        'starting_features': starting_features,
        'talents_by_style_tier': dict(by_tier),
        'recipes': [s['name'] for s in skills if '（配方）' in s['name']],
        'choices': sorted({s['choice'] for s in skills if s.get('choice')}),
        'app': {
            'in_classes_data': '谋士' in app['classes_data'],
            'in_class_features': '谋士' in app['class_features'],
            'in_equipment_data': '谋士' in app['equipment_data'],
            'has_page_json': '谋士' in app['page_json'],
            'has_skill_effects': any('谋士' in x for x in app['skill_effects']),
        },
    }
    report['app_state'] = app
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')

    print('✅ gap report → %s' % OUT)
    print('   谋士技能 %d 条（起始 %d，天赋 %d，配方 %d，抉择组 %d）'
          % (len(skills), len(report['class']['starting_features']),
             sum(len(v) for v in by_tier.values()),
             len(report['class']['recipes']), len(set(report['class']['choices']))))
    print('   谋士进阶 %d 条（全新 %d），需统一来源 %d 条'
          % (len(adv['moushi_cards']), len(adv['moushi_new']), len(adv['unified_sources'])))
    print('   兼职表 %d 职业，谋士兼容 %d，不对称 %d 处'
          % (mc['classes'], len(mc['data'].get('谋士', {}).get('compatible', [])), len(mc['asymmetries'])))
    print('   应用现状：classes_data=%s class_features=%s equipment=%s 页面=%s skill_effects=%s'
          % (report['class']['app']['in_classes_data'], report['class']['app']['in_class_features'],
             report['class']['app']['in_equipment_data'], report['class']['app']['has_page_json'],
             report['class']['app']['has_skill_effects']))
    return 0


if __name__ == '__main__':
    sys.exit(main())
